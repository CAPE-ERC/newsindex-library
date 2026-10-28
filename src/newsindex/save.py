"""Save/load a result table in the format and mechanism a team actually uses.

Consolidates what the brainstorm doc described as two overlapping sections
("Export Helpers": ``to_excel``/``to_json``/``export_all``, and "Saving":
``save``/``save_csv``/.../``save_report``) into a single ``save_*`` family --
``save_all`` already covers what ``export_all`` described, so the former
isn't duplicated here as a separate name.

Every function that needs a library beyond this package's hard dependencies
(``pandas``/``numpy``) -- openpyxl for ``.xlsx``, pyarrow for ``.parquet``,
gspread for Google Sheets, python-docx for ``.docx`` reports, fpdf2 for
``.pdf`` reports -- imports it lazily, inside the function body, and raises
a clear error naming the pip extra to install if it's missing. ``import
newsindex`` itself never requires any of them.
"""

from __future__ import annotations

import datetime
import re
from pathlib import Path

import pandas as pd

# Importing specific names rather than `from . import report as _report`:
# __init__.py does `from .report import report as report`, which rebinds the
# *package's* `report` attribute to that function -- a plain module import
# here would pick up that function instead of the submodule, depending on
# import order. Importing the functions directly sidesteps the collision.
from .report import describe as _describe
from .report import report as _render_report
from .report import summarize as _summarize

_SAVERS_BY_SUFFIX = {
    ".csv": "save_csv",
    ".xlsx": "save_excel",
    ".json": "save_json",
    ".parquet": "save_parquet",
}


def _dispatch_suffix(path: str | Path, table: dict[str, str]) -> str:
    suffix = Path(path).suffix.lower()
    if suffix not in table:
        raise ValueError(
            f"unsupported extension {suffix!r} for {path}; supported: {sorted(table)}"
        )
    return table[suffix]


def save(result: pd.DataFrame, path: str | Path, **kwargs) -> None:
    """Save ``result`` in the format implied by ``path``'s extension.

    ``.csv`` -> :func:`save_csv`, ``.xlsx`` -> :func:`save_excel`, ``.json``
    -> :func:`save_json`, ``.parquet`` -> :func:`save_parquet`. Raises
    :class:`ValueError` for any other extension.
    """
    name = _dispatch_suffix(path, _SAVERS_BY_SUFFIX)
    globals()[name](result, path, **kwargs)


def save_csv(result: pd.DataFrame, path: str | Path, **kwargs) -> None:
    """Save ``result`` as CSV. ``**kwargs`` pass through to ``DataFrame.to_csv``."""
    result.to_csv(path, index=False, **kwargs)


def _sanitize_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\[\]:\*\?/\\]", "_", str(name))
    return cleaned[:31] if cleaned else "Sheet1"


def save_excel(
    result: pd.DataFrame, path: str | Path, *, sheet_by: str | None = None, **kwargs
) -> None:
    """Save ``result`` as a formatted ``.xlsx`` workbook.

    If ``sheet_by`` is given (a column name, e.g. ``"country"``), writes one
    sheet per distinct value of that column; otherwise everything goes on a
    single sheet named ``"result"``. Requires the ``excel`` extra
    (``pip install newsindex[excel]``).
    """
    try:
        writer = pd.ExcelWriter(path, engine="openpyxl")
    except ImportError as exc:
        raise ImportError(
            "save_excel() requires openpyxl -- install with: pip install newsindex[excel]"
        ) from exc

    with writer:
        if sheet_by is None:
            result.to_excel(writer, sheet_name="result", index=False, **kwargs)
        else:
            seen: dict[str, int] = {}
            for value in result[sheet_by].unique():
                sheet_name = _sanitize_sheet_name(value)
                if sheet_name in seen:
                    seen[sheet_name] += 1
                    sheet_name = _sanitize_sheet_name(f"{sheet_name}_{seen[sheet_name]}")
                else:
                    seen[sheet_name] = 0
                result[result[sheet_by] == value].to_excel(
                    writer, sheet_name=sheet_name, index=False, **kwargs
                )


def save_json(result: pd.DataFrame, path: str | Path, **kwargs) -> None:
    """Save ``result`` as tidy records-oriented JSON."""
    result.to_json(path, orient="records", indent=2, **kwargs)


def save_parquet(result: pd.DataFrame, path: str | Path, **kwargs) -> None:
    """Save ``result`` as ``.parquet``. Requires the ``parquet`` extra
    (``pip install newsindex[parquet]``)."""
    try:
        result.to_parquet(path, index=False, **kwargs)
    except ImportError as exc:
        raise ImportError(
            "save_parquet() requires pyarrow -- install with: pip install newsindex[parquet]"
        ) from exc


def save_sql(result: pd.DataFrame, connection, table_name: str, **kwargs) -> None:
    """Write ``result`` straight to a database table.

    ``connection`` is supplied by the caller (e.g. a SQLAlchemy engine or
    connection) -- this function doesn't manage a database driver
    dependency itself, same as ``DataFrame.to_sql`` directly.
    """
    result.to_sql(table_name, connection, index=False, **kwargs)


def save_gsheet(
    result: pd.DataFrame,
    sheet_id_or_url: str,
    *,
    client=None,
    worksheet_index: int = 0,
) -> None:
    """Push ``result`` straight to a Google Sheet, replacing its contents.

    ``client`` is a caller-supplied, already-authorized ``gspread.Client``
    (e.g. from ``gspread.service_account(...)``) -- this function doesn't
    manage Google credentials itself, same caller-supplies-the-connection
    pattern as :func:`save_sql`. If omitted, falls back to
    ``gspread.service_account()`` (default credential discovery) as a
    convenience, but passing ``client`` explicitly is recommended. Requires
    the ``gsheet`` extra (``pip install newsindex[gsheet]``) unless
    ``client`` is already provided by the caller.
    """
    if client is None:
        try:
            import gspread
        except ImportError as exc:
            raise ImportError(
                "save_gsheet() requires gspread -- install with: pip install newsindex[gsheet]"
            ) from exc
        client = gspread.service_account()

    if "http" in sheet_id_or_url:
        sheet = client.open_by_url(sheet_id_or_url)
    else:
        sheet = client.open_by_key(sheet_id_or_url)

    worksheet = sheet.get_worksheet(worksheet_index)
    worksheet.clear()
    worksheet.update([result.columns.tolist()] + result.astype(str).values.tolist())


def save_all(
    results_dict: dict[str, pd.DataFrame], folder: str | Path, formats=("csv",)
) -> list[Path]:
    """Batch-save several result tables in several formats at once.

    ``results_dict`` is ``{name: DataFrame}``. Writes
    ``folder/f"{name}.{fmt}"`` for every ``(name, df)`` and every ``fmt`` in
    ``formats``, creating ``folder`` if needed. Returns the list of paths
    written.
    """
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for name, df in results_dict.items():
        for fmt in formats:
            path = folder / f"{name}.{fmt}"
            save(df, path)
            written.append(path)
    return written


def save_versioned(
    result: pd.DataFrame, folder: str | Path, *, base_name: str = "result", fmt: str = "csv", **kwargs
) -> Path:
    """Timestamped save so a re-run never silently overwrites the last one.

    Writes ``folder/f"{base_name}_{timestamp}.{fmt}"``, creating ``folder``
    if needed, and returns the path written.
    """
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.datetime.now().strftime("%Y%m%dT%H%M%S")
    path = folder / f"{base_name}_{timestamp}.{fmt}"
    save(result, path, **kwargs)
    return path


def autosave(result: pd.DataFrame, project) -> Path:
    """Save ``result`` using a project's configured default path/format.

    ``project`` is a provisional, duck-typed shape pending a real
    ``Project`` config/paths helper class (not implemented yet): either a
    plain ``dict`` with keys ``folder`` (required), ``format`` (default
    ``"csv"``), ``name`` (default ``"result"``), or any object exposing the
    same three as attributes. Saves via :func:`save` to
    ``Path(folder) / f"{name}.{format}"`` and returns that path.
    """
    if isinstance(project, dict):
        folder = project["folder"]
        fmt = project.get("format", "csv")
        name = project.get("name", "result")
    else:
        folder = project.folder
        fmt = getattr(project, "format", "csv")
        name = getattr(project, "name", "result")

    path = Path(folder) / f"{name}.{fmt}"
    save(result, path)
    return path


_LOADERS_BY_SUFFIX = {
    ".csv": pd.read_csv,
    ".xlsx": pd.read_excel,
    ".json": pd.read_json,
    ".parquet": pd.read_parquet,
}


def load(path: str | Path, **kwargs) -> pd.DataFrame:
    """Read back a table saved by :func:`save`, format auto-detected the same way."""
    suffix = Path(path).suffix.lower()
    if suffix not in _LOADERS_BY_SUFFIX:
        raise ValueError(
            f"unsupported extension {suffix!r} for {path}; supported: {sorted(_LOADERS_BY_SUFFIX)}"
        )
    try:
        return _LOADERS_BY_SUFFIX[suffix](path, **kwargs)
    except ImportError as exc:
        extra = {".xlsx": "excel", ".parquet": "parquet"}.get(suffix)
        if extra:
            raise ImportError(
                f"load() of {suffix} files requires an extra package -- "
                f"install with: pip install newsindex[{extra}]"
            ) from exc
        raise


def _docx_report(result: pd.DataFrame, path: str | Path) -> None:
    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            "save_report() with a .docx path requires python-docx -- "
            "install with: pip install newsindex[docx]"
        ) from exc

    document = docx.Document()
    document.add_heading("newsindex EPU report", level=1)
    for line in _summarize(result).splitlines():
        document.add_paragraph(line)

    stats = _describe(result)
    document.add_heading("Summary statistics", level=2)
    table = document.add_table(rows=1, cols=len(stats.columns))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, column in enumerate(stats.columns):
        header_cells[i].text = str(column)
    for row in stats.itertuples(index=False):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = "" if pd.isna(value) else str(value)

    document.save(path)


def _pdf_report(result: pd.DataFrame, path: str | Path) -> None:
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise ImportError(
            "save_report() with a .pdf path requires fpdf2 -- "
            "install with: pip install newsindex[pdf]"
        ) from exc

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "newsindex EPU report", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "", 11)
    for line in _summarize(result).splitlines():
        pdf.multi_cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")

    stats = _describe(result)
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Summary statistics", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "", 9)
    col_width = 190 / max(len(stats.columns), 1)
    for column in stats.columns:
        pdf.cell(col_width, 7, str(column), border=1)
    pdf.ln(7)
    for row in stats.itertuples(index=False):
        for value in row:
            text = "" if pd.isna(value) else str(value)
            pdf.cell(col_width, 7, text[:20], border=1)
        pdf.ln(7)

    pdf.output(str(path))


def save_report(result: pd.DataFrame, path: str | Path, format: str | None = None) -> None:
    """Render a one-page report to ``path``, in more formats than :func:`newsindex.report.report`.

    Format is ``format`` if given, else inferred from ``path``'s extension:
    ``.html``/``.md``/``.markdown`` delegate straight to
    :func:`newsindex.report.report`; ``.docx`` (needs the ``docx`` extra)
    and ``.pdf`` (needs the ``pdf`` extra) build a minimal document here --
    a title, the :func:`newsindex.report.summarize` text, and the
    :func:`newsindex.report.describe` table, same minimalism as
    :func:`newsindex.report.report`. Any other extension raises
    :class:`ValueError`.
    """
    path = Path(path)
    suffix = ("." + format.lstrip(".")) if format else path.suffix.lower()

    if suffix in (".html", ".md", ".markdown"):
        # report() itself picks its format from the path's own extension, so
        # when `format` overrides a path whose extension doesn't match,
        # render to a correctly-suffixed temp file and copy its bytes over
        # instead of duplicating report()'s html/markdown-building logic.
        if path.suffix.lower() == suffix:
            _render_report(result, path)
        else:
            tmp = path.with_name(path.name + suffix)
            _render_report(result, tmp)
            path.write_bytes(tmp.read_bytes())
            tmp.unlink()
    elif suffix == ".docx":
        _docx_report(result, path)
    elif suffix == ".pdf":
        _pdf_report(result, path)
    else:
        raise ValueError(f"unsupported save_report extension {suffix!r}; use .html, .md, .docx, or .pdf")

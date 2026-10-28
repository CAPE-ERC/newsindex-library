"""save.py: save/load a result table in the format and mechanism a team uses."""

import pandas as pd
import pytest

import newsindex as ni

RESULT = pd.DataFrame(
    {
        "country": ["Nigeria", "Nigeria", "Kenya", "Kenya"],
        "year_month": ["2020-01", "2020-02", "2020-01", "2020-02"],
        "EPU": [90.0, 110.0, 80.0, 120.0],
        "M_t": [1.8, 2.2, 1.6, 2.4],
        "mu": [2.0, 2.0, 2.0, 2.0],
        "n_sites": [1, 1, 1, 1],
    }
)


# --- save / load dispatch ---------------------------------------------------

def test_save_and_load_csv_round_trips(tmp_path):
    path = tmp_path / "result.csv"
    ni.save(RESULT, path)
    loaded = ni.load(path)
    pd.testing.assert_frame_equal(loaded, RESULT)


def test_save_unsupported_extension_raises(tmp_path):
    with pytest.raises(ValueError, match=r"unsupported extension"):
        ni.save(RESULT, tmp_path / "result.txt")


def test_load_unsupported_extension_raises(tmp_path):
    with pytest.raises(ValueError, match=r"unsupported extension"):
        ni.load(tmp_path / "result.txt")


def test_save_csv_matches_to_csv(tmp_path):
    path = tmp_path / "result.csv"
    ni.save_csv(RESULT, path)
    pd.testing.assert_frame_equal(pd.read_csv(path), RESULT)


def test_save_json_round_trips_records(tmp_path):
    path = tmp_path / "result.json"
    ni.save_json(RESULT, path)
    loaded = pd.read_json(path)
    assert set(loaded["country"]) == {"Nigeria", "Kenya"}
    assert len(loaded) == len(RESULT)


# --- excel -------------------------------------------------------------------

def test_save_excel_single_sheet(tmp_path):
    pytest.importorskip("openpyxl")
    path = tmp_path / "result.xlsx"
    ni.save_excel(RESULT, path)
    sheets = pd.read_excel(path, sheet_name=None)
    assert list(sheets) == ["result"]
    # check_dtype=False: Excel has no int/float distinction for whole-number
    # values, so e.g. EPU=90.0 round-trips as int64 -- expected Excel
    # behavior, not a save_excel() bug.
    pd.testing.assert_frame_equal(sheets["result"], RESULT, check_dtype=False)


def test_save_excel_sheet_by_splits_sheets(tmp_path):
    pytest.importorskip("openpyxl")
    path = tmp_path / "result.xlsx"
    ni.save_excel(RESULT, path, sheet_by="country")
    sheets = pd.read_excel(path, sheet_name=None)
    assert set(sheets) == {"Nigeria", "Kenya"}
    assert len(sheets["Nigeria"]) == 2
    assert len(sheets["Kenya"]) == 2


def test_load_excel_round_trips(tmp_path):
    pytest.importorskip("openpyxl")
    path = tmp_path / "result.xlsx"
    ni.save_excel(RESULT, path)
    loaded = ni.load(path, sheet_name="result")
    pd.testing.assert_frame_equal(loaded, RESULT, check_dtype=False)


# --- parquet -------------------------------------------------------------------

def test_save_parquet_round_trips(tmp_path):
    pytest.importorskip("pyarrow")
    path = tmp_path / "result.parquet"
    ni.save_parquet(RESULT, path)
    loaded = ni.load(path)
    pd.testing.assert_frame_equal(loaded, RESULT)


# --- sql -----------------------------------------------------------------------

def test_save_sql_writes_to_sqlite(tmp_path):
    sqlalchemy = pytest.importorskip("sqlalchemy")
    engine = sqlalchemy.create_engine(f"sqlite:///{tmp_path / 'db.sqlite'}")
    ni.save_sql(RESULT, engine, "country_index")
    with engine.connect() as conn:
        loaded = pd.read_sql("select * from country_index", conn)
    assert len(loaded) == len(RESULT)
    assert set(loaded["country"]) == {"Nigeria", "Kenya"}


# --- gsheet (fake client, no real gspread/credentials needed) ------------------

class _FakeWorksheet:
    def __init__(self):
        self.cleared = False
        self.updated_with = None

    def clear(self):
        self.cleared = True

    def update(self, rows):
        self.updated_with = rows


class _FakeSheet:
    def __init__(self, worksheet):
        self._worksheet = worksheet

    def get_worksheet(self, index):
        assert index == 0
        return self._worksheet


class _FakeClient:
    def __init__(self, worksheet):
        self._sheet = _FakeSheet(worksheet)
        self.opened_by_key = None
        self.opened_by_url = None

    def open_by_key(self, key):
        self.opened_by_key = key
        return self._sheet

    def open_by_url(self, url):
        self.opened_by_url = url
        return self._sheet


def test_save_gsheet_writes_header_and_rows_via_fake_client():
    worksheet = _FakeWorksheet()
    client = _FakeClient(worksheet)

    ni.save_gsheet(RESULT, "sheet-key-123", client=client)

    assert client.opened_by_key == "sheet-key-123"
    assert worksheet.cleared is True
    assert worksheet.updated_with[0] == list(RESULT.columns)
    assert len(worksheet.updated_with) == len(RESULT) + 1


def test_save_gsheet_uses_open_by_url_for_a_url():
    worksheet = _FakeWorksheet()
    client = _FakeClient(worksheet)

    ni.save_gsheet(RESULT, "https://docs.google.com/spreadsheets/d/abc", client=client)

    assert client.opened_by_url == "https://docs.google.com/spreadsheets/d/abc"


# --- save_all / save_versioned / autosave --------------------------------------

def test_save_all_writes_every_name_and_format(tmp_path):
    written = ni.save_all({"nigeria": RESULT, "kenya": RESULT}, tmp_path, formats=("csv", "json"))
    assert len(written) == 4
    assert (tmp_path / "nigeria.csv").exists()
    assert (tmp_path / "kenya.json").exists()


def test_save_versioned_creates_a_timestamped_file(tmp_path):
    path = ni.save_versioned(RESULT, tmp_path, base_name="country_index")
    assert path.exists()
    assert path.name.startswith("country_index_")
    assert path.suffix == ".csv"


def test_autosave_with_dict_project(tmp_path):
    path = ni.autosave(RESULT, {"folder": str(tmp_path), "format": "csv", "name": "myresult"})
    assert path == tmp_path / "myresult.csv"
    assert path.exists()


class _FakeProject:
    folder = None
    format = "json"
    name = "duck_typed"


def test_autosave_with_duck_typed_project(tmp_path):
    project = _FakeProject()
    project.folder = str(tmp_path)
    path = ni.autosave(RESULT, project)
    assert path == tmp_path / "duck_typed.json"
    assert path.exists()


# --- save_report: docx / pdf -----------------------------------------------------

def test_save_report_docx_is_created_and_nonempty(tmp_path):
    pytest.importorskip("docx")
    path = tmp_path / "report.docx"
    ni.save_report(RESULT, path)
    assert path.exists()
    assert path.stat().st_size > 0

    import docx as docx_module

    document = docx_module.Document(str(path))
    assert len(document.paragraphs) > 0
    assert len(document.tables) == 1


def test_save_report_pdf_is_created_and_nonempty(tmp_path):
    pytest.importorskip("fpdf")
    path = tmp_path / "report.pdf"
    ni.save_report(RESULT, path)
    assert path.exists()
    assert path.stat().st_size > 0
    with open(path, "rb") as handle:
        assert handle.read(4) == b"%PDF"


def test_save_report_unsupported_extension_raises(tmp_path):
    with pytest.raises(ValueError, match=r"unsupported save_report extension"):
        ni.save_report(RESULT, tmp_path / "report.txt")


def test_save_report_format_overrides_extension(tmp_path):
    path = tmp_path / "report.out"
    ni.save_report(RESULT, path, format="html")
    assert path.exists()
    assert "<html>" in path.read_text(encoding="utf-8")
